"""从通讯录 Excel 同步员工到生产/本地数据库。

规则：
- 已存在登录账号（手机号）：更新性别、当前分部、当前部门、角色、登录端口，并重置平台密码为 123456
- 不存在：按表创建员工，密码 123456
- 表内「角色」含「设计师」→ 员工角色「设计师」，登录端口 PC+APP；否则登录端口仅 APP（角色名取表内角色列）

用法：
  docker cp roster.xls api-dev:/tmp/roster.xls
  docker exec api-dev python scripts/sync_employee_roster_xlsx.py /tmp/roster.xls
"""

from __future__ import annotations

import argparse
import asyncio
import re
import sys
import uuid
from pathlib import Path

from dotenv import load_dotenv
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

APP_ROOT = Path(__file__).resolve().parents[1]
for import_path in (APP_ROOT, APP_ROOT / "package"):
    text = str(import_path)
    if text not in sys.path:
        sys.path.insert(0, text)

GENDER_MAP = {"男": "male", "女": "female", "male": "male", "female": "female"}
DESIGNER_ROLE = "设计师"
HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "序号": ("序号",),
    "当前分部": ("当前分部", "分部", "所属门店"),
    "当前部门": ("当前部门", "部门", "所属部门"),
    "员工工号": ("员工工号", "员工编码", "工号", "编号"),
    "员工姓名": ("员工姓名", "姓名", "名字"),
    "手机号码": ("手机号码", "手机号", "登录账号", "手机", "移动电话"),
    "角色": ("角色", "角色名称"),
    "年龄": ("年龄",),
    "员工性别": ("员工性别", "性别"),
}


def load_project_env() -> None:
    load_dotenv(APP_ROOT / ".env", override=False)
    load_dotenv(APP_ROOT.parent / ".env", override=False)


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_phone(value: str) -> str:
    text = _cell_str(value)
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    return text


def _decrypt_if_needed(path: Path, password: str | None) -> Path | None:
    """若文件为加密 OLE，解密到内存；返回 None 表示未加密、直接用原路径。"""
    magic = path.read_bytes()[:4]
    if magic != b"\xd0\xcf\x11\xe0":
        return None
    try:
        import msoffcrypto
    except ImportError as exc:
        raise SystemExit("加密 Excel 需安装 msoffcrypto-tool: uv pip install msoffcrypto-tool xlrd --system") from exc
    import io

    with path.open("rb") as handle:
        office = msoffcrypto.OfficeFile(handle)
        if not office.is_encrypted():
            return None
        if not password:
            raise SystemExit("该 Excel 已加密，请使用 --password 传入打开密码")
        decrypted = io.BytesIO()
        office.load_key(password=password)
        office.decrypt(decrypted)
        decrypted.seek(0)
        return decrypted  # type: ignore[return-value]


def _read_xls(path: Path, *, source: Path | object | None = None) -> list[list[str]]:
    import xlrd

    payload = source if source is not None else path
    if hasattr(payload, "read"):
        book = xlrd.open_workbook(file_contents=payload.read())  # type: ignore[union-attr]
    else:
        book = xlrd.open_workbook(str(payload))
    sheet = book.sheet_by_index(0)
    rows: list[list[str]] = []
    for row_idx in range(sheet.nrows):
        rows.append([_cell_str(sheet.cell_value(row_idx, col)) for col in range(sheet.ncols)])
    return rows


def _read_xlsx(path: Path) -> list[list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    return [[_cell_str(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]


def _map_header(raw_header: list[str]) -> dict[str, int]:
    normalized = {name.strip(): index for index, name in enumerate(raw_header) if name.strip()}
    mapping: dict[str, int] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if alias in normalized:
                mapping[canonical] = normalized[alias]
                break
    required = ("员工工号", "员工姓名", "手机号码", "员工性别", "角色")
    missing = [key for key in required if key not in mapping]
    if missing:
        raise SystemExit(f"表头缺少列: {missing}，实际表头: {raw_header}")
    return mapping


def parse_rows(path: Path, *, password: str | None = None) -> list[dict[str, str]]:
    magic = path.read_bytes()[:4]
    decrypted = _decrypt_if_needed(path, password) if magic == b"\xd0\xcf\x11\xe0" else None
    if magic == b"\xd0\xcf\x11\xe0":
        matrix = _read_xls(path, source=decrypted)
    else:
        matrix = _read_xlsx(path)
    while matrix and not any(matrix[0]):
        matrix.pop(0)
    if not matrix:
        raise SystemExit("表格为空")
    col_map = _map_header(matrix[0])
    rows: list[dict[str, str]] = []
    for raw in matrix[1:]:
        if not raw or not any(raw):
            continue
        record = {key: raw[idx] if idx < len(raw) else "" for key, idx in col_map.items()}
        if not record["员工工号"] or not record["员工姓名"] or not record["手机号码"]:
            continue
        record["手机号码"] = _normalize_phone(record["手机号码"])
        rows.append(record)
    return rows


def resolve_role_and_ports(role_raw: str) -> tuple[str, list[str]]:
    role_text = role_raw.strip()
    if "设计师" in role_text:
        return DESIGNER_ROLE, ["pc", "app"]
    return role_text or "普通用户", ["app"]


async def ensure_role(db, name: str, created_by: str) -> None:
    from yuxi.repositories.role_repository import RoleRepository
    from yuxi.services.role_service import ensure_default_roles, next_role_code

    await ensure_default_roles(db)
    repo = RoleRepository(db)
    if await repo.get_by_name(name):
        return
    await repo.create(
        {
            "id": str(uuid.uuid4()),
            "role_code": next_role_code(await repo.list_codes()),
            "name": name,
            "role_type": "新增",
            "enabled": True,
            "permissions": [],
            "created_by": created_by,
        }
    )


async def reset_platform_password(db, employee) -> None:
    from yuxi.services.employee_service import DEFAULT_EMPLOYEE_PASSWORD, _platform_uid, ensure_platform_user
    from yuxi.storage.postgres.models_business import User
    from yuxi.utils.auth_utils import AuthUtils

    user = (
        await db.execute(select(User).where(User.uid == _platform_uid(employee), User.is_deleted == 0))
    ).scalar_one_or_none()
    if user is None:
        user = await ensure_platform_user(db, employee)
    user.password_hash = AuthUtils.hash_password(DEFAULT_EMPLOYEE_PASSWORD)
    await db.flush()


async def _resolve_creator_uid(db) -> str:
    from yuxi.storage.postgres.models_business import User

    creator = (
        await db.execute(select(User).where(User.role == "superadmin", User.is_deleted == 0).limit(1))
    ).scalar_one_or_none()
    if creator is None:
        creator = (await db.execute(select(User).where(User.is_deleted == 0).limit(1))).scalar_one_or_none()
    if creator is None:
        raise SystemExit("系统中没有可用账号，无法同步员工")
    return str(creator.uid)


async def _sync_one_row(row: dict[str, str], *, creator_uid: str) -> str:
    from yuxi.repositories.employee_repository import EmployeeRepository
    from yuxi.services.employee_service import EmployeeCreate, EmployeeUpdate, create_employee, update_employee
    from yuxi.storage.postgres.manager import pg_manager
    from yuxi.storage.postgres.models_business import User

    gender = GENDER_MAP.get(row["员工性别"].strip())
    if gender is None:
        return "skip"
    age_raw = (row.get("年龄") or "").strip()
    age = int(age_raw) if age_raw.isdigit() else None
    role, login_port = resolve_role_and_ports(row["角色"])
    branch = (row.get("当前分部") or "").strip()
    department = (row.get("当前部门") or "").strip()
    phone = row["手机号码"]
    employee_code = row["员工工号"].strip()

    async with pg_manager.get_async_session_context() as db:
        creator = (
            await db.execute(select(User).where(User.uid == creator_uid, User.is_deleted == 0).limit(1))
        ).scalar_one_or_none()
        if creator is None:
            raise RuntimeError("导入账号不可用")
        await ensure_role(db, role, creator_uid)
        repo = EmployeeRepository(db)
        existing = await repo.get_by_login_account(phone)
        if existing is None:
            existing = await repo.get_by_code(employee_code)

        if existing is not None:
            await update_employee(
                db,
                existing.id,
                EmployeeUpdate(
                    gender=gender,
                    current_branch=branch,
                    current_department=department,
                    role=role,
                    login_port=login_port,
                ),
            )
            employee = await repo.get(existing.id)
            if employee is not None:
                await reset_platform_password(db, employee)
            await db.commit()
            return "updated"

        await create_employee(
            db,
            creator,
            EmployeeCreate(
                employee_code=employee_code,
                name=row["员工姓名"],
                login_account=phone,
                current_branch=branch,
                current_department=department,
                gender=gender,
                age=age,
                login_port=login_port,
                role=role,
            ),
        )
        employee = await repo.get_by_login_account(phone)
        if employee is not None:
            await reset_platform_password(db, employee)
        await db.commit()
        return "created"


async def sync_roster(path: Path, *, password: str | None = None) -> None:
    from yuxi.services.role_service import invalidate_role_member_count_cache
    from yuxi.storage.postgres.manager import pg_manager

    load_project_env()
    rows = parse_rows(path, password=password)
    pg_manager.initialize()
    await pg_manager.create_tables()
    await pg_manager.ensure_business_schema()
    await pg_manager.ensure_content_schema()

    async with pg_manager.get_async_session_context() as db:
        creator_uid = await _resolve_creator_uid(db)

    role_names = {resolve_role_and_ports(row["角色"])[0] for row in rows}
    async with pg_manager.get_async_session_context() as db:
        for role_name in sorted(role_names):
            await ensure_role(db, role_name, creator_uid)
        await db.commit()

    created = updated = skipped = failed = 0
    for row in rows:
        if GENDER_MAP.get(row["员工性别"].strip()) is None:
            print(f"SKIP {row['员工工号']} 未知性别 {row['员工性别']!r}")
            skipped += 1
            continue
        try:
            outcome = await _sync_one_row(row, creator_uid=creator_uid)
            if outcome == "created":
                created += 1
            elif outcome == "updated":
                updated += 1
            else:
                skipped += 1
        except HTTPException as exc:
            failed += 1
            print(f"FAIL {row['员工工号']} {row['员工姓名']}: {exc.detail}")
        except Exception as exc:
            failed += 1
            print(f"FAIL {row['员工工号']} {row['员工姓名']}: {exc}")

    await invalidate_role_member_count_cache()
    await pg_manager.close()
    print(f"同步完成 created={created} updated={updated} skipped={skipped} failed={failed} total={len(rows)}")
    if failed:
        raise SystemExit(1)


def main() -> None:
    parser = argparse.ArgumentParser(description="从通讯录 Excel 同步员工")
    parser.add_argument("path", type=Path, help="Excel 路径 (.xls / .xlsx)")
    parser.add_argument("--password", default=None, help="加密 Excel 的打开密码")
    args = parser.parse_args()
    if not args.path.exists():
        raise SystemExit(f"找不到文件: {args.path}")
    asyncio.run(sync_roster(args.path, password=args.password))


if __name__ == "__main__":
    main()
