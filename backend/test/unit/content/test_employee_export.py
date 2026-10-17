import io

import pytest
from openpyxl import load_workbook

from yuxi.services import employee_service


@pytest.mark.asyncio
async def test_export_preserves_all_rows_display_values_and_text(monkeypatch):
    async def listing(db, keyword=None):
        assert db == "db"
        assert keyword is None
        return {
            "employees": [
                {
                    "employee_code": f"00{index}",
                    "name": "=1+1",
                    "current_branch": "分部",
                    "current_department": "部门",
                    "login_account": "0013800000000",
                    "gender": "female",
                    "age": None,
                    "login_port": ["app", "pc"],
                    "role": "普通用户",
                    "rough_image_count": index,
                    "enterprise_rough_image_count": index + 2,
                    "enabled": index != 0,
                }
                for index in range(25)
            ]
        }

    monkeypatch.setattr(employee_service, "list_employees", listing)
    sheet = load_workbook(io.BytesIO(await employee_service.export_employees("db"))).active
    assert sheet.max_row == 26
    assert list(next(sheet.values)) == [
        "序号",
        "员工编码",
        "姓名",
        "当前分部",
        "当前部门",
        "登录账号",
        "性别",
        "年龄",
        "登录端口",
        "角色",
        "毛坯图上传数（企业共享）",
        "状态",
    ]
    assert list(sheet.values)[1] == (
        1,
        "0024",
        "=1+1",
        "分部",
        "部门",
        "0013800000000",
        "女",
        "-",
        "PC&APP",
        "普通用户",
        26,
        "启用",
    )
    assert sheet["C2"].data_type == "s"
    assert sheet["B2"].data_type == sheet["F2"].data_type == "s"
    assert sheet["A26"].value == 25
    assert sheet["K26"].value == 2
    assert sheet["L26"].value == "禁用"


def _employee(**overrides):
    row = {
        "employee_code": "A",
        "name": "甲",
        "current_branch": "",
        "current_department": "",
        "login_account": "1",
        "gender": "",
        "age": None,
        "login_port": [],
        "role": "普通用户",
        "enterprise_rough_image_count": 0,
        "enabled": True,
    }
    row.update(overrides)
    return row


@pytest.mark.asyncio
async def test_export_groups_same_branch_and_sorts_rough_count_desc(monkeypatch):
    async def listing(db, keyword=None):
        return {
            "employees": [
                _employee(employee_code="A", name="甲", current_branch="河西店", enterprise_rough_image_count=3),
                _employee(employee_code="B", name="乙", current_branch="芙蓉店", enterprise_rough_image_count=9),
                _employee(employee_code="C", name="丙", current_branch="河西店", enterprise_rough_image_count=12),
                _employee(employee_code="D", name="丁", current_branch="", enterprise_rough_image_count=20),
                _employee(employee_code="E", name="戊", current_branch="芙蓉店", enterprise_rough_image_count=1),
            ]
        }

    monkeypatch.setattr(employee_service, "list_employees", listing)
    rows = list(load_workbook(io.BytesIO(await employee_service.export_employees("db"))).active.values)
    assert [(row[3] or "", row[1], row[10]) for row in rows[1:]] == [
        ("河西店", "C", 12),
        ("河西店", "A", 3),
        ("芙蓉店", "B", 9),
        ("芙蓉店", "E", 1),
        ("", "D", 20),
    ]
