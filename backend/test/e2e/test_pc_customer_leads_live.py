"""PC page -> real HTTP routes -> isolated PostgreSQL, including downloaded XLSX."""

import io
import json
import re
import socket
from urllib.parse import urlparse

import pytest
from openpyxl import load_workbook
from patchright.async_api import async_playwright, expect

from test.integration.api.test_case_visit_router import visit_api as visit_api
from yuxi.storage.postgres.models_content import ContentEmployee, ContentMaterialShare


@pytest.mark.e2e
@pytest.mark.asyncio
async def test_pc_browser_search_download_and_admin_scope_with_real_api(visit_api):
    async with visit_api.sessions() as db:
        employee = await db.get(ContentEmployee, "emp-1")
        employee.name, employee.employee_code, employee.current_branch = "员工甲", "000123", "长沙门店"
        employee_without_leads = await db.get(ContentEmployee, "emp-3")
        employee_without_leads.current_branch, employee_without_leads.enabled = "岳阳门店", False
        share = await db.get(ContentMaterialShare, "share-a")
        share.sharer_employee_id, share.sharer_name, share.source_platform = "emp-1", "其他联系人", "pc"
        await db.commit()
    origin = f"http://{socket.gethostbyname('web')}:5173"
    account = 1
    errors = []

    async def respond(route):
        parsed = urlparse(route.request.url)
        if parsed.path == "/api/auth/me":
            return await route.fulfill(
                json={
                    "id": account,
                    "uid": f"visitor-test-user-{account}",
                    "username": f"员工{account}",
                    "role": "admin" if account == 3 else "user",
                    "department_id": 1,
                    "permission_scoped": True,
                    "permissions": ["customer_lead.view_list", "customer_lead.delete"],
                }
            )
        if parsed.path.startswith("/api/customer-leads"):
            response = await visit_api.client.request(
                route.request.method,
                parsed.path + (f"?{parsed.query}" if parsed.query else ""),
                headers=visit_api.tokens[account],
            )
            return await route.fulfill(
                status=response.status_code, headers=dict(response.headers), body=response.content
            )
        return await route.fulfill(json={})

    async with async_playwright() as pw:
        browser = await pw.chromium.launch(
            headless=True,
            args=[
                "--no-sandbox",
                f"--unsafely-treat-insecure-origin-as-secure={origin}",
            ],
        )
        try:
            context = await browser.new_context(viewport={"width": 1440, "height": 960}, accept_downloads=True)
            await context.grant_permissions(["local-network-access"], origin=origin)
            token = visit_api.tokens[1]["Authorization"].removeprefix("Bearer ")
            await context.add_init_script(f"localStorage.setItem('user_token', {json.dumps(token)})")
            await context.route("**/api/**", respond)
            page = await context.new_page()
            page.on("pageerror", lambda error: errors.append(str(error)))
            await page.goto(origin + "/customer-leads", wait_until="domcontentloaded")
            await expect(page.get_by_text("199****0001", exact=True)).to_be_visible(timeout=30000)
            await expect(page.get_by_text("员工甲", exact=True)).to_be_visible()
            await expect(page.get_by_text("其他联系人", exact=True)).to_have_count(0)
            await expect(page.get_by_text("199****0002", exact=True)).to_have_count(0)
            await page.locator(".filters .ant-select").first.click()
            options = page.locator(".ant-select-item-option-content")
            await expect(options.get_by_text("岳阳门店", exact=True)).to_be_visible()
            await expect(options.get_by_text("长沙门店", exact=True)).to_be_visible()
            await options.get_by_text("岳阳门店", exact=True).click()
            await expect(page.get_by_text("暂无客资", exact=True)).to_be_visible()
            await expect(page.locator(".summary-card strong").first).to_have_text("1")
            await page.locator(".filters .ant-select").first.click()
            await expect(options.get_by_text("岳阳门店", exact=True)).to_be_visible()
            await expect(options.get_by_text("长沙门店", exact=True)).to_be_visible()
            await page.keyboard.press("Escape")
            await page.get_by_role("button", name=re.compile(r"重\s*置")).click()
            await expect(page.get_by_text("199****0001", exact=True)).to_be_visible()
            await page.get_by_role("button", name="查看完整号码").click()
            await expect(page.get_by_text("19900000001", exact=True)).to_be_visible()
            await page.get_by_role("button", name="隐藏完整号码").click()
            await page.get_by_placeholder("电话号码、分享人姓名、工号、所属门店").fill("000123")
            await page.get_by_role("button", name="搜索客资", exact=True).click()
            await expect(page.get_by_text("000123", exact=True)).to_be_visible()
            async with page.expect_download() as downloaded:
                await page.get_by_role("button", name="导出", exact=True).click()
            download = await downloaded.value
            sheet = load_workbook(io.BytesIO((await download.path()).read_bytes())).active
            assert sheet.max_row == 2 and sheet["B2"].value == "199****0001"
            assert sheet["E2"].value == "000123" and sheet["F2"].value == "长沙门店"
            await page.screenshot(path="/tmp/pc-customer-leads-live.png", full_page=True)
            await page.evaluate("localStorage.setItem('theme', 'dark')")
            await page.reload(wait_until="domcontentloaded")
            await expect(page.get_by_text("199****0001", exact=True)).to_be_visible()
            await page.screenshot(path="/tmp/pc-customer-leads-dark.png", full_page=True)
            await page.set_viewport_size({"width": 960, "height": 900})
            assert await page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
            account = 3
            await page.reload(wait_until="domcontentloaded")
            await expect(page.get_by_text("暂无客资", exact=True)).to_be_visible()
            await expect(page.locator(".summary-card strong").first).to_have_text("0")
            await page.locator(".filters .ant-select").first.click()
            await expect(options.get_by_text("岳阳门店", exact=True)).to_be_visible()
            await expect(options.get_by_text("长沙门店", exact=True)).to_be_visible()
            await page.keyboard.press("Escape")
            async with page.expect_download() as downloaded:
                await page.get_by_role("button", name="导出", exact=True).click()
            downloaded_path = await (await downloaded.value).path()
            sheet = load_workbook(io.BytesIO(downloaded_path.read_bytes())).active
            assert sheet.max_row == 1
            assert not errors
        finally:
            await browser.close()
