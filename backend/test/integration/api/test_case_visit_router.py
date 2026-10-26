"""Real HTTP and PostgreSQL checks in a disposable schema; no production customers."""

import asyncio
import io
import os
import socket
import uuid
from datetime import datetime, timedelta
from contextlib import asynccontextmanager
import json
from types import SimpleNamespace

import httpx
import pytest
import pytest_asyncio
import uvicorn
from fastapi import FastAPI
from PIL import Image
from openpyxl import load_workbook
from redis.asyncio import Redis
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from server.routers.case_visit_router import case_visits
from server.utils.auth_middleware import get_db
from yuxi.repositories.case_visit_repository import CaseVisitRepository
from yuxi.services import case_visit_service
from yuxi.services.run_queue_service import REDIS_URL
from unittest.mock import AsyncMock
from yuxi.storage.postgres.models_business import Department, User
from yuxi.storage.postgres.models_content import (
    ContentCaseVisitLead,
    ContentCaseVisitSubmission,
    ContentEmployee,
    ContentMaterialShare,
    ContentRole,
)
from yuxi.utils.auth_utils import AuthUtils
from yuxi.utils.datetime_utils import utc_now_naive


@pytest_asyncio.fixture
async def visit_api(monkeypatch):
    schema = f"test_visit_{uuid.uuid4().hex}"
    database_url = os.environ["POSTGRES_URL"].replace("postgresql://", "postgresql+asyncpg://", 1)
    setup_engine = create_async_engine(database_url)
    async with setup_engine.begin() as connection:
        await connection.execute(text(f'CREATE SCHEMA "{schema}"'))
    engine = create_async_engine(database_url, connect_args={"server_settings": {"search_path": schema}})
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        for table in [
            Department.__table__,
            User.__table__,
            ContentRole.__table__,
            ContentEmployee.__table__,
            ContentMaterialShare.__table__,
            ContentCaseVisitLead.__table__,
            ContentCaseVisitSubmission.__table__,
        ]:
            await connection.run_sync(table.create)
    tokens, mp_tokens = {}, {}
    async with sessions() as db:
        db.add_all([Department(id=1, name="部门甲"), Department(id=2, name="部门乙")])
        await db.flush()
        for index, role in enumerate(["user", "user", "admin", "admin", "superadmin"], 1):
            user = User(
                id=index,
                uid=f"visitor-test-user-{index}",
                username=f"visit-{index}",
                role=role,
                department_id=2 if index == 4 else 1,
                password_hash="unused-test-password",
                phone_number=f"1880000000{index}",
            )
            db.add(user)
            tokens[index] = {"Authorization": f"Bearer {AuthUtils.create_access_token({'sub': str(index)})}"}
            db.add(
                ContentEmployee(
                    id=f"emp-{index}",
                    employee_code=f"MP{index}",
                    name=f"员工{index}",
                    login_account=f"1880000000{index}",
                    gender="unknown",
                    role=role,
                    login_port=["app"],
                    created_by=user.uid,
                )
            )
            mp_tokens[index] = {
                "Authorization": f"Bearer {AuthUtils.create_access_token({'sub': f'emp-{index}', 'typ': 'mp'})}"
            }
        db.add_all(
            [
                ContentMaterialShare(
                    id="share-a", token="token-a", owner_uid="material-owner", category_id="case-a", title="原始案例甲"
                ),
                ContentMaterialShare(
                    id="share-b", token="token-b", owner_uid="material-owner", category_id="case-b", title="原始案例乙"
                ),
            ]
        )
        await db.flush()
        data = io.BytesIO()
        Image.new("RGB", (24, 24), "blue").save(data, "WEBP")
        for index, (owner, department, share) in enumerate(
            [(1, 1, "share-a"), (2, 1, "share-b"), (4, 2, "share-a"), (None, None, "share-a")], 1
        ):
            db.add(
                ContentCaseVisitLead(
                    id=f"lead-{index}",
                    share_id=share,
                    appid="wx-test",
                    openid=f"openid-{index}",
                    nickname=f"客户{index}",
                    phone=f"1990000000{index}",
                    avatar_data=data.getvalue(),
                    sharer_uid=f"visitor-test-user-{owner}" if owner else None,
                    department_id=department,
                    created_at=utc_now_naive(),
                    acquired_at=utc_now_naive(),
                )
            )
        await db.commit()

    # Redis 连接跟随测试事件循环，避免跨函数复用已关闭循环上的全局连接。
    redis = Redis.from_url(REDIS_URL, decode_responses=True)
    monkeypatch.setattr(case_visit_service, "get_redis_client", AsyncMock(return_value=redis))
    # 每个独立数据库 schema 使用相同测试员工 ID，清除上一个测试遗留的未读缓存。
    for index in range(1, 6):
        await case_visit_service.invalidate_mp_lead_unread_cache(f"visitor-test-user-{index}")

    async def isolated_db():
        async with sessions() as db:
            yield db

    app = FastAPI()
    app.include_router(case_visits, prefix="/api")
    app.dependency_overrides[get_db] = isolated_db
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    port = sock.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(app, log_level="error", lifespan="off"))
    task = asyncio.create_task(server.serve(sockets=[sock]))
    while not server.started:
        if task.done():
            await task
        await asyncio.sleep(0.01)
    try:
        async with httpx.AsyncClient(base_url=f"http://127.0.0.1:{port}") as client:
            yield SimpleNamespace(client=client, tokens=tokens, mp_tokens=mp_tokens, sessions=sessions)
    finally:
        server.should_exit = True
        await task
        await engine.dispose()
        async with setup_engine.begin() as connection:
            await connection.execute(text(f'DROP SCHEMA "{schema}" CASCADE'))
        await setup_engine.dispose()
        await redis.aclose()


@pytest.mark.asyncio
async def test_real_http_scope_applies_to_listing_avatar_and_delete(visit_api):
    client, tokens = visit_api.client, visit_api.tokens
    assert (await client.get("/api/customer-leads")).status_code == 401
    for user, expected in [
        (1, ["lead-1"]),
        (2, ["lead-2"]),
        (3, []),
        (4, ["lead-3"]),
        (5, ["lead-1", "lead-2", "lead-3", "lead-4"]),
    ]:
        response = await client.get("/api/customer-leads", headers=tokens[user])
        assert response.status_code == 200, response.text
        assert sorted(row["id"] for row in response.json()["items"]) == expected
        for lead in ["lead-1", "lead-2", "lead-3", "lead-4"]:
            avatar = await client.get(f"/api/customer-leads/{lead}/avatar", headers=tokens[user])
            assert avatar.status_code == (200 if lead in expected else 404)
    listing = (
        await client.get(
            "/api/customer-leads",
            headers={
                **tokens[1],
                "X-Forwarded-Proto": "https",
                "X-Forwarded-Host": "cases.test",
                "X-Forwarded-Prefix": "/boyun",
            },
        )
    ).json()
    assert listing["items"][0]["phone_masked"] == "199****0001"
    assert "phone" not in listing["items"][0]
    assert listing["items"][0]["source_url"] == "https://cases.test/boyun/share/case/token-a"
    assert "openid" not in listing["items"][0]
    image = io.BytesIO()
    Image.new("RGB", (16, 16), "blue").save(image, "PNG")
    appointment = await client.post(
        "/api/mp/share/cases/token-a/appointments",
        data={
            "session_id": "expired-test-session",
            "nickname": "测试客户",
            "phone_code": "fake-code",
            "consent": "false",
        },
        files={"avatar": ("avatar.png", image.getvalue(), "image/png")},
    )
    assert appointment.status_code == 422
    assert (await client.get("/api/customer-leads", headers=tokens[5])).json()["total"] == 4
    assert (await client.delete("/api/customer-leads/lead-1", headers=tokens[2])).status_code == 404
    assert (await client.delete("/api/customer-leads/lead-1", headers=tokens[4])).status_code == 404
    assert (await client.delete("/api/customer-leads/lead-1", headers=tokens[1])).status_code == 200
    assert (await client.get("/api/customer-leads/lead-1/avatar", headers=tokens[1])).status_code == 404
    assert (await client.get("/api/customer-leads", headers=tokens[5])).json()["total"] == 3


@pytest.mark.asyncio
async def test_pc_summary_phone_and_export_share_the_same_owner_scope(visit_api):
    client, tokens = visit_api.client, visit_api.tokens
    base = "/api/customer-leads"
    for user, ids in [(1, [1]), (2, [2]), (3, []), (4, [3]), (5, [1, 2, 3, 4])]:
        headers = tokens[user]
        summary = await client.get(base + "/summary", headers=headers)
        assert summary.status_code == 200, summary.text
        assert summary.json()["total"] == len(ids)
        exported = await client.get(base + "/export", headers=headers, params={"page_size": 1})
        assert exported.status_code == 200, exported.text
        assert exported.headers["cache-control"] == "private, no-store"
        sheet = load_workbook(io.BytesIO(exported.content)).active
        assert sheet.max_row == len(ids) + 1
        assert sorted(row[1] for row in list(sheet.values)[1:]) == [f"199****000{i}" for i in ids]
        for index in range(1, 5):
            response = await client.get(f"{base}/lead-{index}/phone", headers=headers)
            assert response.status_code == (200 if index in ids else 404)
            if index in ids:
                assert response.json() == {"phone": f"1990000000{index}"}
                assert response.headers["cache-control"] == "private, no-store"
    for path in ["/summary", "/export", "/lead-1/phone"]:
        assert (await client.get(base + path)).status_code == 401
        assert (await client.get(base + path, headers=visit_api.mp_tokens[1])).status_code == 401
    # A caller-supplied owner or department must never expand access.
    response = await client.get(base, headers=tokens[3], params={"sharer_uid": "visitor-test-user-1"})
    assert response.json()["items"] == []
    assert (await client.delete(base + "/lead-1", headers=tokens[3])).status_code == 404


@pytest.mark.asyncio
async def test_pc_employee_filters_export_and_unfiltered_summary(visit_api):
    base = "/api/customer-leads"
    async with visit_api.sessions() as db:
        employee = await db.get(ContentEmployee, "emp-1")
        employee.name, employee.employee_code, employee.current_branch = "真实员工%甲", "00123", "长沙门店甲"
        share = await db.get(ContentMaterialShare, "share-a")
        share.sharer_employee_id, share.sharer_name, share.source_platform = "emp-1", "手填联系人", "pc"
        lead = await db.get(ContentCaseVisitLead, "lead-1")
        lead.created_at = utc_now_naive() - timedelta(days=2)
        lead.acquired_at = datetime(2026, 10, 15, 16)
        lead2 = await db.get(ContentCaseVisitLead, "lead-2")
        lead2.acquired_at = datetime(2026, 10, 16, 16)
        await db.commit()
    client, headers = visit_api.client, visit_api.tokens[1]
    summary = (await client.get(base + "/summary", headers=headers)).json()
    assert summary == {"total": 1, "today": 0, "branches": ["长沙门店甲"]}
    for keyword in ["19900000001", "真实员工", "00123", "长沙门店", "%"]:
        response = await client.get(base, headers=headers, params={"keyword": keyword})
        assert response.status_code == 200, response.text
        row = response.json()["items"][0]
        assert row["sharer_name"] == "真实员工%甲" and row["employee_code"] == "00123"
        assert row["branch"] == "长沙门店甲" and "phone" not in row
    assert (await client.get(base, headers=headers, params={"keyword": "手填联系人"})).json()["total"] == 0
    for params in [
        {"branch": "其他门店"},
        {"source": "mp"},
        {"keyword": "_"},
        {"start_date": "2026-10-15", "end_date": "2026-10-15"},
    ]:
        assert (await client.get(base, headers=headers, params=params)).json()["total"] == 0
        assert (await client.get(base + "/summary", headers=headers, params=params)).json() == summary
    filters = {"branch": "长沙门店甲", "source": "pc", "start_date": "2026-10-16", "end_date": "2026-10-16"}
    response = await client.get(base, headers=visit_api.tokens[5], params={**filters, "page_size": 1})
    assert response.json()["total"] == 1
    exported = await client.get(base + "/export", headers=visit_api.tokens[5], params=filters)
    sheet = load_workbook(io.BytesIO(exported.content)).active
    assert sheet.max_row == 2 and sheet["E2"].value == "00123" and sheet["E2"].data_type == "s"
    assert sheet["B2"].value.startswith("199****")
    invalid_dates = {"start_date": "2026-10-17", "end_date": "2026-10-16"}
    assert (await client.get(base, headers=headers, params=invalid_dates)).status_code == 422
    assert (await client.get(base + "/export", headers=headers, params={"source": "invalid"})).status_code == 422
    async with visit_api.sessions() as db:
        share = await db.get(ContentMaterialShare, "share-a")
        share.sharer_employee_id = None
        await db.commit()
    # Missing historical employee metadata must not hide leads or change ownership.
    row = (await client.get(base, headers=headers)).json()["items"][0]
    assert row["id"] == "lead-1" and row["sharer_name"] is None and row["employee_code"] is None
    assert (await client.get(base, headers=visit_api.tokens[5])).json()["total"] == 4


@pytest.mark.asyncio
async def test_pc_branch_options_include_all_employee_branches_without_expanding_lead_scope(visit_api):
    async with visit_api.sessions() as db:
        for index, branch in enumerate(["长沙门店", "长沙门店", " 岳阳门店 ", " \t ", ""], 1):
            employee = await db.get(ContentEmployee, f"emp-{index}")
            employee.current_branch = branch
            if index == 3:
                employee.enabled = False
        share = await db.get(ContentMaterialShare, "share-a")
        share.sharer_employee_id = "emp-1"
        await db.commit()

    client, base = visit_api.client, "/api/customer-leads"
    expected_branches = sorted(["长沙门店", "岳阳门店"])
    for user, total in [(1, 1), (2, 1), (3, 0), (4, 1), (5, 4)]:
        headers = visit_api.tokens[user]
        summary = await client.get(base + "/summary", headers=headers)
        assert summary.status_code == 200, summary.text
        assert summary.json()["branches"] == expected_branches
        assert summary.json()["total"] == total
        assert summary.json()["today"] == total

        # The disabled employee has neither shares nor leads, but its branch remains selectable.
        params = {"branch": "岳阳门店"}
        response = await client.get(base, headers=headers, params=params)
        assert response.status_code == 200, response.text
        assert response.json()["items"] == [] and response.json()["total"] == 0
        exported = await client.get(base + "/export", headers=headers, params=params)
        assert exported.status_code == 200, exported.text
        assert load_workbook(io.BytesIO(exported.content)).active.max_row == 1
        assert (await client.get(base + "/summary", headers=headers, params=params)).json() == summary.json()

    # Knowing another employee's branch must not grant access to that employee's leads.
    response = await client.get(base, headers=visit_api.tokens[2], params={"branch": "长沙门店"})
    assert response.json()["items"] == [] and response.json()["total"] == 0


@pytest.mark.asyncio
async def test_postgres_concurrent_upsert_deduplicates_per_share_and_allows_rebooking_after_delete(visit_api):
    async def save(share_id):
        async with visit_api.sessions() as db:
            lead_id = await CaseVisitRepository(db).upsert(
                {
                    "id": uuid.uuid4().hex,
                    "share_id": share_id,
                    "appid": "wx-test",
                    "openid": "same-visitor",
                    "nickname": "新昵称",
                    "phone": "19900000009",
                    "avatar_data": b"updated-avatar",
                    "sharer_uid": "visitor-test-user-1",
                    "department_id": 1,
                    "created_at": utc_now_naive(),
                    "acquired_at": utc_now_naive(),
                }
            )
            await db.commit()
            return lead_id

    ids = await asyncio.gather(*[save("share-a") for _ in range(8)])
    assert len(set(ids)) == 1
    other_id = await save("share-b")
    assert other_id != ids[0]
    async with visit_api.sessions() as db:
        rows = list(
            (await db.scalars(select(ContentCaseVisitLead).where(ContentCaseVisitLead.openid == "same-visitor"))).all()
        )
        assert len(rows) == 2
        assert all(row.nickname == "新昵称" and row.avatar_data == b"updated-avatar" for row in rows)
    assert (
        await visit_api.client.delete(f"/api/customer-leads/{ids[0]}", headers=visit_api.tokens[1])
    ).status_code == 200
    assert await save("share-a") != ids[0]
    async with visit_api.sessions() as db:
        assert (
            await db.scalar(
                select(func.count())
                .select_from(ContentCaseVisitLead)
                .where(ContentCaseVisitLead.openid == "same-visitor")
            )
            == 2
        )


@pytest.mark.asyncio
async def test_mp_owner_scope_including_admin_and_auth_separation(visit_api):
    client = visit_api.client
    base = "/api/mp/customer-leads"
    assert (await client.get(base)).status_code == 401
    assert (await client.get(base, headers=visit_api.tokens[1])).status_code == 401
    assert (await client.get("/api/customer-leads", headers=visit_api.mp_tokens[1])).status_code == 401
    for user, expected in [(1, ["lead-1"]), (2, ["lead-2"]), (3, []), (4, ["lead-3"]), (5, [])]:
        headers = visit_api.mp_tokens[user]
        response = await client.get(base, headers=headers)
        assert response.status_code == 200, response.text
        assert [row["id"] for row in response.json()["items"]] == expected
        assert (await client.get(base + "/summary", headers=headers)).json()["unread_count"] == len(expected)
        for lead in ["lead-1", "lead-2", "lead-3", "lead-4"]:
            expected_status = 200 if lead in expected else 404
            for suffix in ["", "/avatar"]:
                assert (await client.get(f"{base}/{lead}{suffix}", headers=headers)).status_code == expected_status
            if lead not in expected:
                for method, action in [("POST", "read"), ("POST", "follow")]:
                    assert (
                        await client.request(method, f"{base}/{lead}/{action}", headers=headers, json={"revision": 1})
                    ).status_code == 404
    async with visit_api.sessions() as db:
        employee = await db.get(ContentEmployee, "emp-1")
        employee.enabled = False
        await db.commit()
    # Employee enablement is checked on each request, not merely at token issuance.
    assert (await client.get(base, headers=visit_api.mp_tokens[1])).status_code in (401, 403)


@pytest.mark.asyncio
async def test_mp_read_follow_independent_and_old_revision_cannot_consume_new_booking(visit_api):
    client, headers = visit_api.client, visit_api.mp_tokens[1]
    base = "/api/mp/customer-leads"
    detail = (await client.get(base + "/lead-1", headers=headers)).json()
    assert detail["unread"] is True and detail["revision"] == 1
    assert detail["phone"] == "19900000001" and detail["phone_masked"] == "199****0001"
    assert detail["source_share_id"] == "token-a" and detail["source_platform"] is None
    assert (await client.get(base + "/summary", headers=headers)).json()["unread_count"] == 1
    listing = (await client.get(base, headers=headers)).json()["items"][0]
    assert "phone" not in listing and "openid" not in listing
    for _ in range(2):
        result = await client.post(base + "/lead-1/follow", headers=headers, json={"revision": 1})
        assert result.status_code == 200 and result.json()["unread_count"] == 1
    for _ in range(2):
        result = await client.post(base + "/lead-1/read", headers=headers, json={"revision": 1})
        assert result.status_code == 200 and result.json()["unread_count"] == 0
    async with visit_api.sessions() as db:
        lead_id = await CaseVisitRepository(db).upsert(
            {
                "id": "duplicate",
                "share_id": "share-a",
                "appid": "wx-test",
                "openid": "openid-1",
                "nickname": "重新预约",
                "phone": "19900000001",
                "avatar_data": b"avatar",
                "sharer_uid": "visitor-test-user-1",
                "acquired_at": utc_now_naive(),
            }
        )
        await db.commit()
        assert lead_id == "lead-1"
    # 直接写库模拟预约时，同步执行预约服务提交后的缓存失效动作。
    await case_visit_service.invalidate_mp_lead_unread_cache("visitor-test-user-1")
    for method, action in [("POST", "read"), ("POST", "follow")]:
        assert (
            await client.request(method, base + "/lead-1/" + action, headers=headers, json={"revision": 1})
        ).status_code == 409
    updated = (await client.get(base + "/lead-1", headers=headers)).json()
    assert updated["revision"] == 2 and updated["unread"] is True and updated["follow_status"] == "pending"
    assert (await client.get(base + "/summary", headers=headers)).json()["unread_count"] == 1
    assert (await client.post(base + "/lead-1/read", headers=headers, json={"revision": 2})).status_code == 200
    assert (await client.get(base + "/lead-1", headers=headers)).json()["follow_status"] == "pending"


@pytest.mark.asyncio
async def test_mp_search_pagination_and_beijing_date_boundaries(visit_api):
    async with visit_api.sessions() as db:
        lead = await db.get(ContentCaseVisitLead, "lead-1")
        lead.nickname = "客户%甲"
        lead.acquired_at = datetime(2026, 10, 15, 16, 0)
        db.add(
            ContentCaseVisitLead(
                id="boundary",
                share_id="share-b",
                appid="wx-test",
                openid="boundary",
                nickname="客户乙",
                phone="19911112222",
                avatar_data=b"a",
                sharer_uid=lead.sharer_uid,
                acquired_at=datetime(2026, 10, 16, 16, 0),
            )
        )
        await db.commit()
    client, headers = visit_api.client, visit_api.mp_tokens[1]
    base = "/api/mp/customer-leads"
    for params, expected in [
        ({"keyword": "%"}, ["lead-1"]),
        ({"keyword": "1111"}, ["boundary"]),
        ({"start_date": "2026-10-16", "end_date": "2026-10-16"}, ["lead-1"]),
        ({"page_size": 1, "page": 1}, ["boundary"]),
        ({"page_size": 1, "page": 2}, ["lead-1"]),
    ]:
        result = await client.get(base, headers=headers, params=params)
        assert result.status_code == 200, result.text
        assert [row["id"] for row in result.json()["items"]] == expected
    assert (
        await client.get(base, headers=headers, params={"start_date": "2026-10-17", "end_date": "2026-10-16"})
    ).status_code == 422


@pytest.mark.asyncio
async def test_http_booking_receipt_survives_retry_and_new_session_rebooks(visit_api, monkeypatch):
    class Sessions:
        def __init__(self):
            self.data = {
                f"mp:visit:{name}": json.dumps({"appid": "wx-test", "openid": "visitor-retry", "share_token": token})
                for name, token in [("one", "token-a"), ("two", "token-a"), ("three", "token-b")]
            }
            self.mutex = asyncio.Lock()
            self.delete = AsyncMock()

        async def get(self, key):
            return self.data.get(key)

        async def setex(self, key, ttl, value):
            self.data[key] = value

        @asynccontextmanager
        async def lock(self, *args, **kwargs):
            async with self.mutex:
                yield

    phone = AsyncMock(return_value="19912345678")
    monkeypatch.setattr(case_visit_service, "get_redis_client", AsyncMock(return_value=Sessions()))
    monkeypatch.setattr(case_visit_service, "_wechat_credentials", lambda: ("wx-test", "unused"))
    monkeypatch.setattr(case_visit_service, "_wechat_phone_by_code", phone)
    async with visit_api.sessions() as db:
        for share_id in ["share-a", "share-b"]:
            share = await db.get(ContentMaterialShare, share_id)
            share.sharer_uid = "visitor-test-user-1"
            share.source_platform = "mp"
            share.sharer_phone = "19912345678" if share_id == "share-a" else "19900000002"
        await db.commit()
    image = io.BytesIO()
    Image.new("RGB", (24, 24), "red").save(image, "PNG")

    async def book(session, token="token-a", *, profile=True):
        data = {"session_id": session, "phone_code": "verified-code", "consent": "true"}
        if profile:
            data["nickname"] = "预约客户"
        response = await visit_api.client.post(
            f"/api/mp/share/cases/{token}/appointments",
            data=data,
            files={"avatar": ("avatar.png", image.getvalue(), "image/png")} if profile else None,
        )
        assert response.status_code == 200, response.text
        assert response.json()["same_contact_phone"] is (token == "token-a")
        return response.json()["id"]

    ids = await asyncio.gather(book("one"), book("one"))
    assert ids[0] == ids[1]
    phone.assert_awaited_once()
    base = f"/api/mp/customer-leads/{ids[0]}"
    headers = visit_api.mp_tokens[1]
    assert (await visit_api.client.post(base + "/read", headers=headers, json={"revision": 1})).status_code == 200
    assert (await visit_api.client.post(base + "/follow", headers=headers, json={"revision": 1})).status_code == 200
    assert await book("one") == ids[0]
    saved = (await visit_api.client.get(base, headers=headers)).json()
    assert saved["revision"] == 1 and saved["unread"] is False and saved["follow_status"] == "followed"
    assert saved["source_platform"] == "mp"
    assert await book("two", profile=False) == ids[0]
    saved = (await visit_api.client.get(base, headers=headers)).json()
    assert saved["revision"] == 2 and saved["unread"] is True and saved["follow_status"] == "pending"
    assert saved["nickname"] == "预约客户" and saved["avatar_url"]
    assert (await visit_api.client.get(saved["avatar_url"], headers=headers)).status_code == 200
    new_id = await book("three", "token-b", profile=False)
    assert new_id != ids[0]
    new_lead = (await visit_api.client.get(f"/api/mp/customer-leads/{new_id}", headers=headers)).json()
    assert new_lead["nickname"] == "预约客户" and new_lead["avatar_url"] is None
    assert new_lead["phone"] == "19912345678" and new_lead["unread"]
    pc_list = (await visit_api.client.get("/api/customer-leads", headers=visit_api.tokens[1])).json()
    pc_lead = next(item for item in pc_list["items"] if item["id"] == new_id)
    assert not {"nickname", "avatar_url", "phone"}.intersection(pc_lead)
    assert pc_lead["phone_masked"] == "199****5678"
    async with visit_api.sessions() as db:
        stored = await db.get(ContentCaseVisitLead, new_id)
        assert stored.nickname is None and stored.avatar_data is None
    refused = await visit_api.client.post(
        "/api/mp/share/cases/token-b/appointments",
        data={"session_id": "three", "phone_code": "new-code", "consent": "false"},
    )
    assert refused.status_code == 422
    assert phone.await_count == 3
