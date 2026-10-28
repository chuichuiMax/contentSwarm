"""Real HTTP and PostgreSQL report/deletion checks in a disposable schema."""

import asyncio
import io
from datetime import datetime, timedelta
from unittest.mock import AsyncMock

import pytest
import pytest_asyncio
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import func, select

from server.routers.case_share_router import case_shares
from server.routers.material_library_router import material_library, public_share_router
from test.integration.api.test_case_visit_router import visit_api as visit_api
from yuxi.repositories.case_visit_repository import CaseVisitRepository
from yuxi.repositories.material_library_repository import MaterialLibraryRepository
from yuxi.services import case_share_service as stats
from yuxi.services import case_visit_service as visits
from yuxi.services import material_share_images as images
from yuxi.services import material_share_jobs as jobs
from yuxi.storage.postgres.models_business import OperationLog, User
from yuxi.storage.postgres.models_content import (
    ContentCaseVisitLead,
    ContentCaseVisitSubmission,
    ContentEmployee,
    ContentMaterialShare,
    ContentMaterialShareItem,
    ContentMaterialShareJob,
    ContentRole,
)
from yuxi.repositories.case_share_repository import CaseShareRepository
from yuxi.utils.datetime_utils import SHANGHAI_TZ

pytestmark = [pytest.mark.integration, pytest.mark.asyncio]
TOKEN = "a" * 64
START = datetime(2026, 10, 27, 16)


@pytest_asyncio.fixture
async def report(visit_api, monkeypatch):
    monkeypatch.setattr(stats, "shanghai_now", lambda: datetime(2026, 10, 28, 12, tzinfo=SHANGHAI_TZ))
    monkeypatch.setattr(images, "get_binary_redis_client", visits.get_redis_client)
    visit_api.app.include_router(case_shares, prefix="/api")
    visit_api.app.include_router(material_library, prefix="/api")
    visit_api.app.include_router(public_share_router)
    async with visit_api.sessions() as db:
        async with db.bind.begin() as conn:
            for model in (OperationLog, ContentMaterialShareJob, ContentMaterialShareItem):
                await conn.run_sync(model.__table__.create)
        first = await db.get(ContentMaterialShare, "share-a")
        first.token = TOKEN
        first.source_platform = "mp"
        first.sharer_uid = "visitor-test-user-1"
        first.sharer_employee_id = "emp-1"
        first.sharer_employee_name = "实际员工"
        first.sharer_employee_code = "0001"
        first.sharer_branch = "分享时门店"
        first.sharer_name = "代接联系人"
        first.mp_entry_code = "c" * 24
        first.created_at = START
        second = await db.get(ContentMaterialShare, "share-b")
        second.source_platform = "pc"
        second.sharer_uid = "visitor-test-user-2"
        second.sharer_employee_name = "员工乙"
        second.sharer_employee_code = "0002"
        second.sharer_branch = "乙门店"
        second.created_at = START - timedelta(microseconds=1)
        db.add(
            ContentMaterialShare(
                id="historical",
                token="legacy",
                owner_uid="visitor-test-user-1",
                category_id="case-a",
                title="历史案例",
                created_at=START + timedelta(days=1),
            )
        )
        db.add(
            ContentMaterialShareJob(
                token=TOKEN,
                requester_uid=first.sharer_uid,
                item_ids=["image"],
                share_json={"title": "原始案例甲"},
                sources_json=[],
                status="completed",
            )
        )
        db.add(
            ContentMaterialShareItem(
                share_id=first.id,
                display_order=1,
                original_file_name="snapshot.png",
                content_type="image/png",
                file_size=4,
                image_width=1,
                image_height=1,
                bucket_name="image",
                object_name="share-only/snapshot.png",
            )
        )
        db.add(ContentCaseVisitSubmission(session_id="receipt", lead_id="lead-1"))
        for state, token in (("pending", "b" * 64), ("failed", "d" * 64)):
            db.add(
                ContentMaterialShareJob(
                    token=token,
                    requester_uid=first.sharer_uid,
                    item_ids=["image"],
                    share_json={},
                    sources_json=[],
                    status=state,
                )
            )
        await db.commit()
    yield visit_api


async def test_report_scope_snapshot_pagination_and_links(report):
    client, base = report.client, "/api/case-shares"
    assert (await client.get(base)).status_code == 401
    expected = {1: ["share-a"], 2: ["share-b"], 3: [], 4: [], 5: ["historical", "share-a", "share-b"]}
    for user, ids in expected.items():
        response = await client.get(base, headers=report.tokens[user])
        assert response.status_code == 200, response.text
        assert [row["id"] for row in response.json()["items"]] == ids
        assert response.headers["cache-control"] == "private, no-store"
    own = (await client.get(base, headers=report.tokens[1])).json()["items"][0]
    assert (own["sharer_name"], own["employee_code"], own["branch"]) == ("实际员工", "0001", "分享时门店")
    assert own["source_url"].endswith(f"/share/case/{TOKEN}")
    assert "phone" not in own and "title" not in own
    prefixed = await client.get(
        base,
        headers={
            **report.tokens[1],
            "X-Forwarded-Prefix": "/boyun",
            "X-Forwarded-Host": "share.example.test",
            "X-Forwarded-Proto": "https",
        },
    )
    assert prefixed.json()["items"][0]["source_url"] == f"https://share.example.test/boyun/share/case/{TOKEN}"
    async with report.sessions() as db:
        employee = await db.get(ContentEmployee, "emp-1")
        employee.name, employee.employee_code, employee.current_branch = "后来姓名", "9999", "后来门店"
        await db.commit()
    assert (await client.get(base, headers=report.tokens[1])).json()["items"][0] == own
    page = await client.get(base, headers=report.tokens[5], params={"page": 2, "page_size": 1})
    assert page.json()["total"] == 3 and page.json()["items"][0]["id"] == "share-a"
    historical = (await client.get(base, headers=report.tokens[5])).json()["items"][0]
    assert historical["sharer_name"] is None and historical["branch"] is None and historical["source_platform"] is None


async def test_filters_fixed_summary_and_beijing_boundaries(report):
    client, base, headers = report.client, "/api/case-shares", report.tokens[5]
    for keyword, count in [
        ("0001", 1),
        ("实际员工", 1),
        ("分享时门店", 1),
        ("代接联系人", 0),
        ("原始案例甲", 0),
        ("%", 0),
    ]:
        response = await client.get(base, headers=headers, params={"keyword": keyword})
        assert response.json()["total"] == count
    for source in ("pc", "mp"):
        assert (await client.get(base, headers=headers, params={"source": source})).json()["total"] == 1
    dates = {"start_date": "2026-10-28", "end_date": "2026-10-28"}
    rows = (await client.get(base, headers=headers, params=dates)).json()
    assert rows["total"] == 1 and rows["items"][0]["id"] == "share-a"
    assert (await client.get(base + "/summary", headers=headers, params={"keyword": "no-match", **dates})).json() == {
        "total": 3,
        "today": 1,
    }
    assert (await client.get(base + "/summary", headers=report.tokens[1])).json() == {"total": 1, "today": 1}
    assert (await client.get(base + "/summary", headers=report.tokens[3])).json() == {"total": 0, "today": 0}
    for params in ({"source": "unknown"}, {"page_size": 101}, {"start_date": "2026-10-29", "end_date": "2026-10-28"}):
        assert (await client.get(base, headers=headers, params=params)).status_code == 422


async def test_export_all_filtered_rows_text_cells_and_permissions(report):
    client, base, headers = report.client, "/api/case-shares", report.tokens[5]
    async with report.sessions() as db:
        row = await db.get(ContentMaterialShare, "share-a")
        row.sharer_employee_name = "=1+1"
        await db.commit()
    response = await client.get(base + "/export", headers=headers, params={"source": "mp"})
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.max_row == 2
    assert [cell.value for cell in sheet[2]] == [1, "小程序案例分享", "=1+1", "0001", "分享时门店", "2026-10-28 00:00"]
    assert sheet["C2"].data_type == "s" and sheet["D2"].data_type == "s"
    assert (
        load_workbook(io.BytesIO((await client.get(base + "/export", headers=report.tokens[3])).content)).active.max_row
        == 1
    )
    async with report.sessions() as db:
        db.add(
            ContentRole(
                id="report-role",
                role_code="report-role",
                name="无报表权限",
                created_by="test",
                permissions=["chat.access"],
            )
        )
        user = await db.get(User, 1)
        user.role = "无报表权限"
        await db.commit()
    for suffix in ("", "/summary", "/export"):
        assert (await client.get(base + suffix, headers=report.tokens[1])).status_code == 403
    assert (await client.delete(base + "/share-a", headers=report.tokens[1])).status_code == 403


async def test_view_only_cannot_delete_and_non_superadmin_cannot_delete_others(report):
    client, base = report.client, "/api/case-shares"
    for user in (2, 3, 4):
        assert (await client.delete(base + "/share-a", headers=report.tokens[user])).status_code == 404
    async with report.sessions() as db:
        db.add(
            ContentRole(
                id="read-only",
                role_code="read-only",
                name="只读报表",
                created_by="test",
                permissions=["customer_lead.view_list"],
            )
        )
        user = await db.get(User, 1)
        user.role = "只读报表"
        await db.commit()
    assert (await client.get(base, headers=report.tokens[1])).json()["total"] == 1
    assert (await client.get(base + "/export", headers=report.tokens[1])).status_code == 200
    assert (await client.delete(base + "/share-a", headers=report.tokens[1])).status_code == 403


async def test_delete_removes_appointments_receipts_caches_and_all_public_entries(report):
    redis = await visits.get_redis_client()
    key = f"material-share:v1:{TOKEN}:cover"
    await redis.set(key, "cached-image")
    await visits.set_mp_lead_unread_cache("visitor-test-user-1", 99)
    response = await report.client.delete("/api/case-shares/share-a", headers=report.tokens[1])
    assert response.status_code == 200, response.text
    assert await redis.get(key) is None
    assert await redis.get(visits.mp_lead_unread_cache_key("visitor-test-user-1")) is None
    async with report.sessions() as db:
        assert await db.get(ContentMaterialShare, "share-a") is None
        assert (
            await db.scalar(
                select(func.count()).select_from(ContentCaseVisitLead).where(ContentCaseVisitLead.share_id == "share-a")
            )
            == 0
        )
        assert await db.get(ContentCaseVisitLead, "lead-2") is not None
        assert await db.get(ContentCaseVisitSubmission, "receipt") is None
        assert await db.scalar(select(func.count()).select_from(ContentMaterialShareItem)) == 0
        assert (await db.get(ContentMaterialShareJob, TOKEN)).status == "deleted"
        assert await db.get(ContentMaterialShare, "share-b") is not None
        log = await db.scalar(select(OperationLog).where(OperationLog.operation == "material.share.delete"))
        assert log.user_id == 1
        # The same accepted request stays terminal; the worker cannot recreate it.
        with pytest.raises(HTTPException) as error:
            await jobs.accept_material_share(
                db, await db.get(User, 1), jobs.MaterialShareRequest(item_ids=["image"], request_token=TOKEN), None
            )
        assert error.value.status_code == 410
    monkeypatch_worker = AsyncMock()
    # A deleted job is skipped before any snapshot/storage work.
    from unittest.mock import patch

    with (
        patch.object(jobs.pg_manager, "get_async_session_context", report.sessions),
        patch.object(jobs, "store_material_share_snapshot", monkeypatch_worker),
    ):
        await jobs.process_material_share({}, TOKEN)
    monkeypatch_worker.assert_not_awaited()
    for suffix in ("", "/page", "/booking", "/mp-code", "/open-mp", "/images/1", "/images/1.webp", "/cover.jpg"):
        response = await report.client.get(f"/api/material-library/shares/{TOKEN}{suffix}")
        assert response.status_code in (404, 410), (suffix, response.text)
    assert (await report.client.get(f"/share/case/{TOKEN}")).status_code == 410
    assert (await report.client.get("/api/material-library/share-entries/" + "c" * 24)).status_code == 404
    assert (
        await report.client.post("/api/mp/visitor-sessions", json={"share_token": TOKEN, "code": "unused"})
    ).status_code == 404
    assert (await report.client.get("/api/case-shares/summary", headers=report.tokens[5])).json() == {
        "total": 2,
        "today": 0,
    }


async def test_booking_lock_serializes_share_deletion(report):
    async with report.sessions() as booking:
        assert await CaseVisitRepository(booking).share(TOKEN, for_update=True) is not None
        deletion = asyncio.create_task(report.client.delete("/api/case-shares/share-a", headers=report.tokens[1]))
        try:
            await asyncio.sleep(0.2)
            assert not deletion.done()
            # A booking that committed first must also be removed by the deletion.
            await CaseVisitRepository(booking).upsert(
                {
                    "id": "concurrent",
                    "share_id": "share-a",
                    "appid": "wx-test",
                    "openid": "concurrent",
                    "phone": "19900000009",
                    "sharer_uid": "visitor-test-user-1",
                }
            )
            await booking.commit()
            assert (await deletion).status_code == 200
        finally:
            await booking.rollback()
            await deletion
    async with report.sessions() as db:
        assert await CaseVisitRepository(db).share(TOKEN, for_update=True) is None
        assert await db.get(ContentCaseVisitLead, "concurrent") is None


async def test_generated_case_stats_decrease_without_a_separate_counter(report):
    async with report.sessions() as db:
        repo = MaterialLibraryRepository(db)
        assert await repo.count_case_shares_by_owner(["visitor-test-user-1"]) == {"visitor-test-user-1": 2}
    await report.client.delete("/api/case-shares/share-a", headers=report.tokens[5])
    async with report.sessions() as db:
        assert await MaterialLibraryRepository(db).count_case_shares_by_owner(["visitor-test-user-1"]) == {
            "visitor-test-user-1": 1
        }


async def test_failed_delete_rolls_back_share_leads_receipts_and_job(report, monkeypatch):
    original_delete = CaseShareRepository.delete

    async def failing_delete(repo, share_id, predicates):
        await original_delete(repo, share_id, predicates)
        raise RuntimeError("test rollback after all delete statements")

    monkeypatch.setattr(CaseShareRepository, "delete", failing_delete)
    response = await report.client.delete("/api/case-shares/share-a", headers=report.tokens[1])
    assert response.status_code == 500
    async with report.sessions() as db:
        assert await db.get(ContentMaterialShare, "share-a") is not None
        assert await db.get(ContentCaseVisitLead, "lead-1") is not None
        assert await db.get(ContentCaseVisitSubmission, "receipt") is not None
        assert (await db.get(ContentMaterialShareJob, TOKEN)).status == "completed"


async def test_legacy_share_without_job_can_be_deleted_by_superadmin_only(report):
    assert (await report.client.delete("/api/case-shares/historical", headers=report.tokens[1])).status_code == 404
    response = await report.client.delete("/api/case-shares/historical", headers=report.tokens[5])
    assert response.status_code == 200, response.text
    assert (await report.client.get("/api/material-library/shares/legacy")).status_code == 404
